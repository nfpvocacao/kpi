import os
import sys
import time
import datetime
import openpyxl
import threading
import pyautogui
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager

def disparar_confirmacao_certificado():
    """
    Thread auxiliar para aguardar a janela modal de Certificado Digital do Windows
    e pressionar ENTER para selecionar o certificado 'AÇÃO COMUNITÁRIA DO BRASIL'.
    """
    print("  [Certificado] Aguardando prompt do Windows para selecionar certificado...")
    time.sleep(3)
    pyautogui.press('enter')
    print("  [Certificado] Tecla ENTER enviada com sucesso!")

def extrair_dados_nfp():
    """
    Acessa o portal NFP com certificado digital, navega até a consulta de doadores
    e extrai as informações das duas tabelas.
    """
    print("\n--- INICIANDO ACESSO AO PORTAL NFP ---")
    url_login = "https://www.nfp.fazenda.sp.gov.br/login.aspx?ReturnUrl=/"
    
    options = webdriver.ChromeOptions()
    options.add_argument("--start-maximized")
    options.add_argument("--ignore-certificate-errors")
    options.add_argument("--disable-popup-blocking")
    
    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=options)
    
    dados_extraidos = {}

    try:
        print(f"1. Acessando URL: {url_login}")
        driver.get(url_login)
        
        wait = WebDriverWait(driver, 20)
        
        # Seleciona Contribuinte ICMS
        print("2. Selecionando opção Contribuinte ICMS...")
        rd_contribuinte = wait.until(EC.element_to_be_clickable((By.ID, "rdBtnContribuinte")))
        rd_contribuinte.click()
        time.sleep(1)
        
        # Dispara thread para lidar com a janela nativa do Windows do certificado
        thread_cert = threading.Thread(target=disparar_confirmacao_certificado)
        thread_cert.start()
        
        # Clica no botão de Certificado Digital
        print("3. Clicando em 'Acesso via certificado digital'...")
        btn_cert = wait.until(EC.element_to_be_clickable((By.ID, "imgBtnAcessoCertCNPJ")))
        btn_cert.click()
        
        thread_cert.join()
        
        print("4. Aguardando autenticação e carregamento da área logada...")
        time.sleep(5)
        
        # Navegação pelos menus: Consultar -> Entidades -> Consulta de Consumidores Doadores
        print("5. Navegando pelo menu: Consultar -> Entidades -> Consulta de Consumidores Doadores...")
        menu_consultar = wait.until(EC.element_to_be_clickable((By.XPATH, "//a[contains(text(),'Consultar')] | //span[contains(text(),'Consultar')]")))
        menu_consultar.click()
        time.sleep(1)
        
        menu_entidades = wait.until(EC.element_to_be_clickable((By.XPATH, "//a[contains(text(),'Entidades')] | //span[contains(text(),'Entidades')]")))
        menu_entidades.click()
        time.sleep(1)
        
        menu_doadores = wait.until(EC.element_to_be_clickable((By.XPATH, "//a[contains(text(),'Consulta de Consumidores Doadores')]")))
        menu_doadores.click()
        
        print("6. Aguardando 15 segundos para renderização completa das tabelas...")
        time.sleep(15)
        
        print("7. Extraindo tabelas da página...")
        
        # Mapeamento de meses abreviados em PT para número do mês
        meses_map = {
            'jan': 1, 'fev': 2, 'mar': 3, 'abr': 4, 'mai': 5, 'jun': 6,
            'jul': 7, 'ago': 8, 'set': 9, 'out': 10, 'nov': 11, 'dez': 12
        }

        def extrair_numero(texto):
            if not texto: return 0
            # Remove R$, pontos de milhar e substitui vírgula decimal
            limpo = str(texto).replace('R$', '').replace('.', '').replace(',', '.').strip()
            try:
                return float(limpo) if '.' in limpo else int(limpo)
            except:
                return 0

        # --- IMPRESSÃO / DEBUG DA ESTRUTURA DAS TABELAS ---
        html_page = driver.page_source
        with open("temp_nfp_page.html", "w", encoding="utf-8") as f:
            f.write(html_page)
        print("  [DEBUG] Conteúdo HTML salvo em temp_nfp_page.html")
        
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html_page, 'html.parser')
        
        tabelas = soup.find_all('table')
        print(f"  Total de tabelas encontradas no DOM: {len(tabelas)}")
        
        for idx, t in enumerate(tabelas):
            texto_tabela = t.get_text(separator=' | ', strip=True)
            print(f"\n  --- TABELA {idx+1} ---")
            print(texto_tabela[:300]) # imprime os primeiros 300 caracteres de cada tabela

        # Processamento flexível para capturar qualquer célula contendo mês (ex: jul/2026)
        for t in tabelas:
            rows = t.find_all('tr')
            for r in rows:
                cols = [c.get_text(strip=True) for c in r.find_all(['td', 'th'])]
                if not cols: continue
                
                # Se for uma linha da tabela de doadores futuros (ex: nov/2026 | 4396 | 826 | 5222)
                if len(cols) >= 4 and '/20' in cols[0]:
                    try:
                        m_str, a_str = cols[0].lower().split('/')
                        m_num = meses_map.get(m_str[:3], 1)
                        dt_key = f"{int(a_str):04d}-{m_num:02d}-01"
                        
                        if dt_key not in dados_extraidos:
                            dados_extraidos[dt_key] = {}
                            
                        dados_extraidos[dt_key]['plenos'] = extrair_numero(cols[1])
                        dados_extraidos[dt_key]['restritos'] = extrair_numero(cols[2])
                        dados_extraidos[dt_key]['total'] = extrair_numero(cols[3])
                        dados_extraidos[dt_key]['cupons'] = 0
                        dados_extraidos[dt_key]['valor_nf'] = 0.0
                    except Exception as err:
                        pass
                
                # Se for a linha de meses da Tabela Situação Atual
                elif any('/20' in c for c in cols):
                    colunas_meses_situacao = []
                    for idx_c, cell in enumerate(cols):
                        if '/20' in cell:
                            try:
                                m_str, a_str = cell.lower().split('/')
                                m_num = meses_map.get(m_str[:3], 1)
                                dt_key = f"{int(a_str):04d}-{m_num:02d}-01"
                                colunas_meses_situacao.append((idx_c, dt_key))
                                if dt_key not in dados_extraidos:
                                    dados_extraidos[dt_key] = {}
                            except:
                                pass
                                
                    # Se encontrou meses nesta tabela, lê as linhas seguintes
                    for r_sub in rows:
                        cols_sub = [c.get_text(strip=True) for c in r_sub.find_all(['td', 'th'])]
                        if not cols_sub: continue
                        label = cols_sub[0].lower()
                        for idx_c, dt_key in colunas_meses_situacao:
                            if idx_c < len(cols_sub):
                                val = extrair_numero(cols_sub[idx_c])
                                if 'plenos' in label:
                                    dados_extraidos[dt_key]['plenos'] = val
                                elif 'restritos' in label:
                                    dados_extraidos[dt_key]['restritos'] = val
                                elif 'valor' in label or 'val' in label:
                                    dados_extraidos[dt_key]['valor_nf'] = val
                                elif 'cupons' in label:
                                    dados_extraidos[dt_key]['cupons'] = val
                                elif 'total doadores' in label or label == 'total':
                                    dados_extraidos[dt_key]['total'] = val

        # --- FILTRAGEM RIGOROSA DA REGRA DINÂMICA (Mês mais antigo + 4 subsequentes = 5 meses) ---
        todos_meses_ordenados = sorted(list(dados_extraidos.keys()))
        
        if todos_meses_ordenados:
            # O primeiro mês é o mês mais antigo (ex: 2026-07-01)
            mes_mais_antigo = todos_meses_ordenados[0]
            
            # Pega exatamente o mês mais antigo + os 4 meses subsequentes (5 meses no total)
            idx_inicio = todos_meses_ordenados.index(mes_mais_antigo)
            meses_permitidos = set(todos_meses_ordenados[idx_inicio : idx_inicio + 5])
            
            # Descarta qualquer mês fora da janela exata dos 5 meses
            dados_filtrados = {k: v for k, v in dados_extraidos.items() if k in meses_permitidos}
            dados_extraidos = dados_filtrados


        print(f"\n  Total de meses filtrados pela regra (Mês mais antigo + 4 subsequentes): {len(dados_extraidos)}")
        for k, v in sorted(dados_extraidos.items()):
            print(f"    - {k}: {v}")

    except Exception as e:
        print(f"ERRO DURANTE O ACESSO/EXTRAÇÃO NO PORTAL NFP: {e}")
    finally:
        driver.quit()
        
    return dados_extraidos



def atualizar_planilha_mapa(caminho_excel, dados):
    """
    Atualiza as colunas da aba 'Mapa' na planilha Excel local com os dados extraídos.
    Colunas: Doadores Plenos, Doadores Restritos, Total Doadores, Qtde Cupons, Val NF.
    """
    if not os.path.exists(caminho_excel):
        print(f"Erro: Planilha não encontrada em: {caminho_excel}")
        return False

    print(f"\n--- ATUALIZANDO PLANILHA EXCEL LOCAL ({caminho_excel}) ---")
    wb = openpyxl.load_workbook(caminho_excel)
    
    if 'Mapa' not in wb.sheetnames:
        print("Erro: Aba 'Mapa' não encontrada na planilha!")
        return False
        
    ws = wb['Mapa']
    
    col_mes = 1           # Coluna A (Mês)
    col_plenos = 2        # Coluna B (Doadores Plenos)
    col_restritos = 4     # Coluna D (Doadores Restritos)
    col_total = 5         # Coluna E (Total Doadores)
    col_cupons = 6        # Coluna F (Qtde Cupons)
    col_val_nf = 7        # Coluna G (Val NF)
    
    atualizados = 0
    
    for row in range(2, ws.max_row + 1):
        cell_mes = ws.cell(row=row, column=col_mes).value
        if cell_mes is None:
            continue
            
        dt_str = None
        if isinstance(cell_mes, datetime.datetime) or isinstance(cell_mes, datetime.date):
            dt_str = cell_mes.strftime("%Y-%m-%d")
        else:
            dt_str = str(cell_mes).strip()[:10]
            
        if dt_str in dados:
            item = dados[dt_str]
            print(f"Atualizando linha {row} ({dt_str}): Plenos={item.get('plenos')}, Restritos={item.get('restritos')}, Total={item.get('total')}, Cupons={item.get('cupons')}, ValNF={item.get('valor_nf')}")
            
            if 'plenos' in item:
                ws.cell(row=row, column=col_plenos, value=item['plenos'])
            if 'restritos' in item:
                ws.cell(row=row, column=col_restritos, value=item['restritos'])
            if 'total' in item:
                ws.cell(row=row, column=col_total, value=item['total'])
            if 'cupons' in item:
                ws.cell(row=row, column=col_cupons, value=item['cupons'])
            if 'valor_nf' in item:
                ws.cell(row=row, column=col_val_nf, value=item['valor_nf'])
                
            atualizados += 1

    wb.save(caminho_excel)
    print(f"Planilha atualizada e salva com sucesso! Total de meses atualizados: {atualizados}")
    return True

if __name__ == "__main__":
    caminho_mapa = r"G:\Drives compartilhados\NFP\Fazenda\Automatizações\2025-01-01 Mapa Automatizacoes.xlsx"
    print(f"Iniciando automação NFP. Alvo da planilha: {caminho_mapa}")
    
    dados = extrair_dados_nfp()
    if dados:
        atualizar_planilha_mapa(caminho_mapa, dados)
    else:
        print("Nenhum dado retornado da extração.")


